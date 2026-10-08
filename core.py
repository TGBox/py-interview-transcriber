"""Reine Logik: Sprecherzuordnung, Absätze, Ausgabeformen, HTML/DOCX. Nur stdlib (+ python-docx lazy)."""
import re
from html import escape

Segment = tuple[float, float, str]  # (start, end, text) – Whisper-Wort oder -Segment
Turn = tuple[float, float, str]  # (start, end, speaker) aus pyannote
Block = tuple  # ("h1"|"h2"|"li"|"text", text) oder ("p", zeit|None, sprecher, text)

PAUSE = "(...)"  # Pausenmarke nach Dresing/Pehl
PAUSE_SECONDS = 3.0
FORMS = {
    "woertlich": "Wörtlich",
    "geglaettet": "Geglättet",
    "geglaettet_llm": "Geglättet (Sprachmodell)",
    "wissenschaftlich": "Wissenschaftlich (Dresing/Pehl)",
    "zusammenfassung": "Sinngemäße Zusammenfassung",
}

_FILLER = re.compile(r"(?i)(?:,\s*)?\b(?:äh+m*|öh+m*|ehm|hm+)\b,?")  # bewusst ohne "mhm" (oft = Ja) und "eh" (Wort)
# Artikel/Pronomen ausgenommen: "die die", "Sie sie", "das das" sind im Deutschen oft korrekt
_STUTTER = re.compile(r"(?i)\b(?!(?:der|die|das|den|dem|des|sie|ihr|ihre|was|wer|ob)\b)(\w+)(?:\s+\1\b)+")
# Großschreiben nur nach echtem Satzende: nicht nach "z. B.", "ca.", "bzw.", "usw.", "etc."
_SENTENCE_START = re.compile(r"(^|(?<=\w{3})(?<!bzw)(?<!usw)(?<!etc)[.?!]\s+)(\w)")


def assign_speakers(segments: list[Segment], turns: list[Turn]) -> list[dict]:
    """Jedem Segment/Wort den Sprecher mit der größten zeitlichen Überlappung geben."""
    out = []
    for start, end, text in segments:
        speaker = "SPEAKER_00"
        if turns:
            # ponytail: O(wörter × turns) – bei 1 h Interview ~10k × 500, unkritisch; Sweep-Line erst bei Bedarf
            overlap = lambda t: min(end, t[1]) - max(start, t[0])
            best = max(turns, key=overlap)
            if overlap(best) <= 0:  # liegt in einer Pause: nächster Turn
                best = min(turns, key=lambda t: max(t[0] - end, start - t[1]))
            speaker = best[2]
        out.append({"start": start, "end": end, "speaker": speaker, "text": text})
    return out


TERMINAL_PUNCT = re.compile(r'[.?!][\"\'»”’“\)]?\s*$')


def split_first_sentence(text: str) -> tuple[str, str]:
    """Splits text into (first_sentence, remainder)."""
    text = text.strip()
    m = re.match(r'^\s*([^.?!]+[.?!]+[\"\'»”’“\)]?)(?:\s+(.*))?$', text, re.DOTALL)
    if m:
        return m.group(1).strip(), (m.group(2) or "").strip()
    return text, ""


def split_last_sentence(text: str) -> tuple[str, str]:
    """Splits text into (remainder_before_last_sentence, last_sentence)."""
    text = text.strip()
    m = re.match(r'^(.*[.?!]+[\"\'»”’“\)]?)\s+([^.?!]+[.?!]+[\"\'»”’“\)]?)$', text, re.DOTALL)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    m2 = re.match(r'^(.*[.?!]+[\"\'»”’“\)]?)\s+([^.?!]+)$', text, re.DOTALL)
    if m2:
        return m2.group(1).strip(), m2.group(2).strip()
    return "", text


def split_trailing_fragment(text: str, max_words: int = 4) -> tuple[str, str | None]:
    """Checks if text ends with terminal punctuation followed by a short trailing fragment."""
    text = text.strip()
    m = re.match(r'^(.*[.?!]+[\"\'»”’“\)]?)\s+([^.?!]+)$', text, re.DOTALL)
    if m:
        trailing = m.group(2).strip()
        if len(trailing.split()) <= max_words:
            return m.group(1).strip(), trailing
    return text, None


def realign_paragraph_boundaries(paragraphs: list[dict]) -> tuple[list[dict], int]:
    """Realigns misattributed sentence ends/beginnings at speaker boundaries."""
    out = [dict(p) for p in paragraphs]
    changes = 0
    i = 0
    while i < len(out) - 1:
        p0 = out[i]
        p1 = out[i + 1]
        t0 = p0.get("text", "").strip()
        t1 = p1.get("text", "").strip()

        # Case A: p0 ends with complete sentence + trailing fragment (<= 4 words) and p1 starts lowercase
        base0, trailing = split_trailing_fragment(t0)
        if trailing and t1 and t1[0].islower():
            p0["text"] = base0
            p1["text"] = trailing + " " + t1
            dur = p0["end"] - p0["start"]
            frac = len(trailing) / max(len(t0), 1)
            time_shift = dur * frac
            p0["end"] = max(p0["start"], p0["end"] - time_shift)
            p1["start"] = min(p1["end"], p1["start"] - time_shift)
            p0.pop("smooth", None)
            p1.pop("smooth", None)
            changes += 1
            i += 1
            continue

        # Case B: p0 does not end in terminal punct, and p1 starts with continuation fragment ending in punct
        if not TERMINAL_PUNCT.search(t0):
            first_sent, rest = split_first_sentence(t1)
            if first_sent and (t1[0].islower() or len(first_sent.split()) <= 6):
                p0["text"] = t0 + " " + first_sent
                dur = p1["end"] - p1["start"]
                frac = len(first_sent) / max(len(t1), 1)
                time_shift = dur * frac
                p0["end"] = min(p1["end"], p0["end"] + time_shift)
                p0.pop("smooth", None)
                p1.pop("smooth", None)
                if rest:
                    p1["text"] = rest
                    p1["start"] = min(p1["end"], p1["start"] + time_shift)
                    i += 1
                else:
                    # p1 completely absorbed into p0
                    p0["end"] = p1["end"]
                    out.pop(i + 1)
                changes += 1
                continue

        i += 1

    return out, changes


def merge_paragraphs(items: list[dict], pause: float = PAUSE_SECONDS) -> list[dict]:
    """Aufeinanderfolgende Wörter desselben Sprechers zu Absätzen zusammenfassen, lange Pausen markieren."""
    out: list[dict] = []
    for it in items:
        if out and out[-1]["speaker"] == it["speaker"]:
            gap = it["start"] - out[-1]["end"]
            out[-1]["text"] += f" {PAUSE} " + it["text"] if gap >= pause else " " + it["text"]
            out[-1]["end"] = it["end"]
        else:
            out.append(dict(it))
    realigned, _ = realign_paragraph_boundaries(out)
    return realigned


def clean_text(text: str, keep_pauses: bool = False) -> str:
    """Geglättet: Füllwörter, Stottern und (optional) Pausenmarken entfernen."""
    if not keep_pauses:
        text = text.replace(PAUSE, " ")
    text = _FILLER.sub("", text)
    text = _STUTTER.sub(r"\1", text)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([,.?!])", r"\1", text).strip().lstrip(",.;: ")
    return _SENTENCE_START.sub(lambda m: m.group(1) + m.group(2).upper(), text)


def fmt_time(seconds: float, tenths: bool = False) -> str:
    t = int(seconds * 10 + 1e-6)
    s = t // 10
    hms = f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"
    return f"{hms}-{t % 10}" if tenths else hms


def speaker_display_name(label: str) -> str:
    """Wandelt technische Bezeichner wie SPEAKER_00 in sprechende deutsche Namen um."""
    m = re.match(r"^(?:SPEAKER|SPK)[_ -]?(\d+)$", label, re.IGNORECASE)
    if m:
        return f"Sprecher {int(m.group(1)) + 1}"
    return label


def render(title: str, paragraphs: list[dict], names: dict[str, str], form: str) -> list[Block]:
    """Absätze in Blöcke der gewählten Form umwandeln (nicht für 'zusammenfassung')."""
    blocks: list[Block] = [("h1", title)]
    for p in paragraphs:
        name = names.get(p["speaker"], speaker_display_name(p["speaker"]))
        if form == "woertlich":
            blocks.append(("p", fmt_time(p["start"]), name, p["text"]))
        elif form == "geglaettet":
            if text := clean_text(p["text"]):
                blocks.append(("p", fmt_time(p["start"]), name, text))
        elif form == "geglaettet_llm":  # Spalte aus der KI-Glättung; reine Füllwort-Absätze entfallen
            if clean_text(p["text"]):
                blocks.append(("p", fmt_time(p["start"]), name, p["smooth"]))
        elif form == "wissenschaftlich":
            # Dresing/Pehl (einfach): geglättet, Pausen (...), Zeitmarke am Absatzende
            if text := clean_text(p["text"], keep_pauses=True):
                blocks.append(("p", None, name, f"{text} #{fmt_time(p['end'], tenths=True)}#"))
        else:
            raise ValueError(form)
    return blocks


def summary_blocks(title: str, markdown: str) -> list[Block]:
    """Das kleine Markdown-Subset der LLM-Antwort (## / - / *) in Blöcke wandeln."""
    blocks: list[Block] = [("h1", title)]
    for line in markdown.splitlines():
        line = line.strip().replace("**", "")
        if line.startswith("#"):
            blocks.append(("h2", line.lstrip("# ")))
        elif line[:2] in ("- ", "* "):
            blocks.append(("li", line[2:]))
        elif line:
            blocks.append(("text", line))
    return blocks


SPEAKER_PALETTE = [
    "#1d4ed8",  # Blau
    "#7e22ce",  # Violett
    "#047857",  # Smaragdgrün
    "#c2410c",  # Rostorange
    "#0891b2",  # Cyan
    "#b91c1c",  # Dunkelrot
]


def to_html(blocks: list[Block]) -> str:
    title = ""
    rows = []
    speaker_colors: dict[str, str] = {}

    def get_speaker_color(speaker_name: str) -> str:
        if speaker_name not in speaker_colors:
            speaker_colors[speaker_name] = SPEAKER_PALETTE[len(speaker_colors) % len(SPEAKER_PALETTE)]
        return speaker_colors[speaker_name]

    for b in blocks:
        kind = b[0]
        if kind == "h1":
            title = b[1]
        elif kind == "h2":
            rows.append(
                f'<tr>'
                f'<td colspan="2" style="padding: 16pt 0 6pt 0; border-bottom: 1.5px solid #cbd5e1;">'
                f'<span style="font-size: 13pt; font-weight: bold; color: #0f172a;">{escape(b[1])}</span>'
                f'</td>'
                f'</tr>'
            )
        elif kind == "p":
            ts = b[1] or ""
            speaker = b[2]
            text = b[3]
            color = get_speaker_color(speaker)
            ts_html = f'<div style="color: #64748b; font-size: 8pt; margin-top: 2pt; font-family: monospace;">[{escape(ts)}]</div>' if ts else ""
            rows.append(
                f'<tr>'
                f'<td width="24%" style="vertical-align: top; padding: 7pt 6pt 7pt 0; border-bottom: 1px solid #e2e8f0;">'
                f'<span style="font-size: 9.5pt; font-weight: bold; color: {color};">{escape(speaker)}</span>'
                f'{ts_html}'
                f'</td>'
                f'<td width="76%" style="vertical-align: top; padding: 7pt 0 7pt 8pt; border-bottom: 1px solid #e2e8f0; font-size: 10pt; line-height: 1.55; color: #1e293b;">'
                f'{escape(text)}'
                f'</td>'
                f'</tr>'
            )
        elif kind == "li":
            rows.append(
                f'<tr>'
                f'<td colspan="2" style="padding: 3pt 0 3pt 14pt; border-bottom: none; font-size: 10pt; color: #334155;">'
                f'• {escape(b[1])}'
                f'</td>'
                f'</tr>'
            )
        elif kind == "text":
            rows.append(
                f'<tr>'
                f'<td colspan="2" style="padding: 6pt 0; border-bottom: none; font-size: 10pt; line-height: 1.5; color: #334155;">'
                f'{escape(b[1])}'
                f'</td>'
                f'</tr>'
            )
        else:
            rows.append(
                f'<tr>'
                f'<td colspan="2" style="padding: 4pt 0;">'
                f'<{kind}>{escape(b[1])}</{kind}>'
                f'</td>'
                f'</tr>'
            )

    title_html = (
        f'<table width="100%" cellpadding="0" cellspacing="0" style="margin-bottom: 14pt; border-bottom: 2px solid #0284c7; padding-bottom: 8pt;">'
        f'<tr>'
        f'<td>'
        f'<div style="font-size: 8.5pt; font-weight: bold; color: #0284c7; letter-spacing: 1px;">INTERVIEW-TRANSKRIPT</div>'
        f'<div style="font-size: 18pt; font-weight: bold; color: #0f172a; margin-top: 3pt;">{escape(title)}</div>'
        f'</td>'
        f'</tr>'
        f'</table>'
    ) if title else ""

    return (
        '<!DOCTYPE html>'
        '<html>'
        '<head>'
        '<meta charset="utf-8">'
        '<style>'
        'body { font-family: "Segoe UI", Arial, Helvetica, sans-serif; font-size: 10pt; color: #1e293b; }'
        '</style>'
        '</head>'
        '<body>'
        f'{title_html}'
        f'<table width="100%" cellpadding="0" cellspacing="0">'
        f'{"".join(rows)}'
        f'</table>'
        '</body>'
        '</html>'
    )


def to_docx(path: str, blocks: list[Block], line_numbers: bool = False) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    for b in blocks:
        kind = b[0]
        if kind == "p":
            para = doc.add_paragraph()
            if b[1]:
                ts = para.add_run(f"[{b[1]}] ")
                ts.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                ts.font.size = Pt(9)
            para.add_run(b[2] + ": ").bold = True
            para.add_run(b[3])
        elif kind == "li":
            doc.add_paragraph(b[1], style="List Bullet")
        elif kind == "text":
            doc.add_paragraph(b[1])
        else:
            doc.add_heading(b[1], level=1 if kind == "h1" else 2)
    if line_numbers:  # Zeilennummern wie in wissenschaftlichen Transkripten üblich
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sect = doc.sections[0]._sectPr
        cols = sect.find(qn("w:cols"))  # Schema-Reihenfolge: lnNumType steht vor cols
        cols.addprevious(ln) if cols is not None else sect.append(ln)
    doc.save(path)
